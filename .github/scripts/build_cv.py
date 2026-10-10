from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path("cv")
OUT.mkdir(exist_ok=True)

CONTACT="Cairo, Egypt | (+20) 1200031122 | khedr6333@gmail.com | linkedin.com/in/abdulrahman-khedr | github.com/Abdulrahman6333 | abdulrahman6333.github.io/AbdulrahmanKhedr/"
BLUE=RGBColor(31,78,121)
DARK=RGBColor(35,35,35)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=45, start=55, bottom=45, end=55):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m,v in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node=OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def add_heading(doc, text):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text)
    r.bold=True; r.font.size=Pt(10); r.font.color.rgb=BLUE
    return p

def add_bullets(doc, bullets):
    for b in bullets:
        p=doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent=Inches(0.18)
        p.paragraph_format.first_line_indent=Inches(-0.1)
        p.paragraph_format.space_after=Pt(1)
        r=p.add_run(b); r.font.size=Pt(8.3)

def setup():
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(0.34); sec.bottom_margin=Inches(0.34)
    sec.left_margin=Inches(0.42); sec.right_margin=Inches(0.42)
    styles=doc.styles
    styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(8.6)
    styles['Normal'].paragraph_format.space_after=Pt(1)
    styles['List Bullet'].font.name='Arial'; styles['List Bullet'].font.size=Pt(8.3)
    return doc

def header(doc, title):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(0)
    r=p.add_run("ABDULRAHMAN KHEDR"); r.bold=True; r.font.name="Arial"; r.font.size=Pt(18); r.font.color.rgb=DARK
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    r=p.add_run(title); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=BLUE
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
    r=p.add_run(CONTACT); r.font.size=Pt(7.4)

def skills_table(doc, rows):
    t=doc.add_table(rows=0, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(1.45); t.columns[1].width=Inches(5.7)
    for a,b in rows:
        cells=t.add_row().cells
        cells[0].width=Inches(1.45); cells[1].width=Inches(5.7)
        cells[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        for c in cells: set_cell_margins(c)
        p=cells[0].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        rr=p.add_run(a); rr.bold=True; rr.font.size=Pt(8.2); rr.font.color.rgb=DARK
        p=cells[1].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        rr=p.add_run(b); rr.font.size=Pt(8.2)

def add_role(doc, line, bullets=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(line); r.bold=True; r.font.size=Pt(8.7)
    if bullets: add_bullets(doc, bullets)

def save(doc, name):
    path=OUT/name
    doc.save(path)
    return path

def build_python():
    d=setup(); header(d,"Python Applications & AI Automation Developer | Mechatronics Engineer")
    add_heading(d,"PROFESSIONAL SUMMARY")
    p=d.add_paragraph("Mechatronics Engineer and Teaching Assistant with hands-on experience building Python desktop applications, academic decision-support tools, and AI-powered automation workflows. Strong in Python, n8n, REST APIs, SQLite, Streamlit, PySide6, Selenium, Excel/PDF automation, Telegram integrations, OpenRouter, and structured AI-agent workflows. Combines software development with an engineering background in embedded systems, robotics, and control.")
    p.paragraph_format.space_after=Pt(2)
    add_heading(d,"CORE TECHNICAL SKILLS")
    skills_table(d,[
        ("Python & Apps","Python, Streamlit, Tkinter, PySide6, Pandas, SQLite"),
        ("Automation & AI","n8n, AI Agents, OpenRouter, REST APIs, structured JSON workflows, prompt design"),
        ("Data & Documents","OpenPyXL, Excel automation, PyMuPDF, ReportLab, PDF generation"),
        ("Integrations","Telegram bots, Selenium, SMTP/email automation, FFmpeg, TTS"),
        ("Engineering","C/C++, AVR, STM32, Arduino, Raspberry Pi, ROS, MATLAB/Simulink"),
        ("Tools","Git, GitHub, Ubuntu/Linux, SOLIDWORKS"),
    ])
    add_heading(d,"SELECTED SOFTWARE & AUTOMATION PROJECTS")
    add_role(d,"Graduation Planner & Academic Advising System | Python • Streamlit • Pandas • Excel/PDF",[
        "Built an academic decision-support application that analyzes student history and supports structured graduation planning.",
        "Implemented GPA/CGPA logic, failed-course and retake tracking, prerequisites, electives, remaining-course analysis, semester planning, bilingual UI, Excel export, and portrait PDF reports."
    ])
    add_role(d,"AI E-commerce Automation & Product Intelligence | n8n • OpenRouter • AI Agents • Telegram • FFmpeg",[
        "Designed workflows for AI-assisted product research, evidence-aware ranking, structured output, Telegram interaction, and short-form content automation.",
        "Built a pipeline covering product selection, real-image intake, creative planning, Arabic TTS, FFmpeg assembly, and delivery while avoiding fabricated product claims."
    ])
    add_role(d,"Energix Communication Automation Suite | Python • Tkinter • Selenium • SMTP • Pillow",[
        "Developed an Excel-driven desktop suite for WhatsApp messaging, personalized email delivery, certificate generation, reusable templates, and PDF attachments."
    ])
    add_role(d,"Engineering Course File Manager | Python • PySide6 • SQLite",[
        "Built a privacy-safe portfolio edition for course management, synthetic faculty assignments, an 18-item course-file checklist, progress tracking, and reporting."
    ])
    add_heading(d,"EXPERIENCE")
    add_role(d,"Teaching Assistant — Delta University for Science & Technology | Mar 2024 – Present",[
        "Teach and support mechatronics courses and labs covering microcontrollers, PLC/SCADA, robotics, control, sensors, circuits, and computer vision.",
        "Supervise student projects and develop Python-based tools and automation for academic and administrative workflows."
    ])
    add_role(d,"IMT Egypt — AVR Embedded Systems Diploma | Jul 2022 – Nov 2022",["Embedded C, AVR microcontrollers, peripheral interfacing, and hardware-software integration."])
    add_role(d,"Smart Technology Egypt — Engineering Training | Aug 2021 – Oct 2021",["Practical training in Arduino-based systems and automation."])
    add_heading(d,"EDUCATION")
    add_role(d,"Pre-Master’s in Mechatronics Engineering — Mansoura University | Oct 2024 – May 2025 | Grade: A")
    add_role(d,"B.Sc. in Mechatronics Engineering — Higher Technological Institute | Aug 2018 – Jun 2023 | Grade: Very Good (B+)",[
        "Graduation Project: 6-Axis Industrial Robot Arm — Grade A+; MATLAB/ROS simulation and Arduino/custom-driver implementation; adopted by ITAC (Smart Village)."
    ])
    add_heading(d,"LANGUAGES")
    d.add_paragraph("Arabic — Native | English — Professional Working Proficiency")
    return save(d,"Abdulrahman_Khedr_CV_Python_AI_Automation.docx")

def build_mech():
    d=setup(); header(d,"Mechatronics Engineer | Embedded Systems • Robotics • Control")
    add_heading(d,"PROFESSIONAL SUMMARY")
    d.add_paragraph("Mechatronics Engineer and Teaching Assistant with practical experience in embedded systems, robotics, control, microcontrollers, industrial automation, and engineering software. Skilled in C/C++, Python, AVR, STM32, Arduino, Raspberry Pi, ROS, MATLAB/Simulink, PLC/SCADA, communication protocols, and hardware-software integration. Also develops Python tools and automation systems that support engineering and academic workflows.")
    add_heading(d,"CORE TECHNICAL SKILLS")
    skills_table(d,[
        ("Embedded Systems","C, C++, AVR, STM32, Arduino, Raspberry Pi, Embedded Linux"),
        ("Robotics & Control","ROS, MATLAB, Simulink, robot kinematics, path planning, control systems"),
        ("Industrial Automation","PLC, SCADA, sensors, actuators, instrumentation concepts"),
        ("Protocols","UART, SPI, I2C, CAN"),
        ("Software","Python, PySide6, Streamlit, SQLite, Git/GitHub"),
        ("Mechanical Design","SOLIDWORKS, engineering design and prototyping"),
    ])
    add_heading(d,"ENGINEERING PROJECTS")
    add_role(d,"6-Axis Industrial Robot Arm | Graduation Project | Grade A+",[
        "Designed and simulated a six-axis industrial robot-arm system using MATLAB and ROS, including control, inverse kinematics, and path-planning workflows.",
        "Implemented hardware control using Arduino and custom motor-driver interfacing; project was adopted by ITAC (Smart Village)."
    ])
    add_role(d,"Real-Time Powertrain Control / Embedded Linux | Embedded & Control Project",[
        "Worked on Linux-based control and real-time engineering concepts relevant to embedded mechatronic systems and motor-control applications."
    ])
    add_role(d,"Engineering Course File Manager | Python • PySide6 • SQLite",[
        "Developed a desktop engineering workflow application demonstrating software architecture, database use, UI design, and requirements-driven engineering."
    ])
    add_role(d,"Graduation Planner & Academic Advising System | Python • Streamlit",[
        "Built a rules-based decision-support application with data processing, validation logic, reporting, and bilingual engineering-software workflows."
    ])
    add_role(d,"Energix Communication Automation Suite | Python Desktop Application",[
        "Built an automation application integrating Excel data, browser automation, SMTP email, certificate generation, and PDF handling."
    ])
    add_heading(d,"EXPERIENCE")
    add_role(d,"Teaching Assistant — Delta University for Science & Technology | Mar 2024 – Present",[
        "Teach and support laboratories and courses in microcontrollers, PLC/SCADA, robotics, control systems, sensors, circuits, and computer vision.",
        "Supervise engineering projects, prepare laboratory material, support academic advising, and contribute to lab development and technical troubleshooting."
    ])
    add_role(d,"IMT Egypt — AVR Embedded Systems Diploma | Jul 2022 – Nov 2022",["Embedded C, AVR architecture, timers, ADC, interrupts, peripheral interfacing, and hardware-software integration."])
    add_role(d,"Smart Technology Egypt — Engineering Training | Aug 2021 – Oct 2021",["Practical training in Arduino-based systems and automation."])
    add_heading(d,"EDUCATION")
    add_role(d,"Pre-Master’s in Mechatronics Engineering — Mansoura University | Oct 2024 – May 2025 | Grade: A")
    add_role(d,"B.Sc. in Mechatronics Engineering — Higher Technological Institute | Aug 2018 – Jun 2023 | Grade: Very Good (B+)")
    add_heading(d,"SOFT SKILLS & LANGUAGES")
    d.add_paragraph("Problem Solving • Technical Communication • Teamwork • Leadership • Self-learning • Requirements Analysis")
    d.add_paragraph("Arabic — Native | English — Professional Working Proficiency")
    return save(d,"Abdulrahman_Khedr_CV_Mechatronics_Embedded.docx")

if __name__=="__main__":
    build_python(); build_mech()
